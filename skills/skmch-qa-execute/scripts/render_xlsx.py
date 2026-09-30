#!/usr/bin/env python3
"""Fill an Excel (.xlsx) template from JSON rows.

JSON input shape:
{
  "meta":   {"DOC_ID": "CR-2026-014", "VERSION": "0.1", ...},   # fills {{KEY}} in any cell
  "sheets": {
    "Test Cases": [ {"TC ID": "TC-001", "Covers": "FR-001", ...}, ... ],
    "Traceability": [ ... ]
  }
}

For each sheet the script finds the header row (the first row with a cell that equals
a key of the first JSON row) and writes the rows beneath it, matching columns by header
text (case-insensitive). The first data row's formatting is copied if it exists.
A sheet missing from the template is created, with headers taken from the JSON keys.

Usage: python render_xlsx.py data.json out.xlsx [--template templates/x.xlsx]
Needs openpyxl (pip install openpyxl).
"""
import argparse
import copy
import datetime
import json
import re
import sys

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

TOKEN = re.compile(r"\{\{([A-Z0-9_]+)\}\}")
HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")


def norm(v):
    return str(v).strip().lower() if v is not None else ""


def find_header(ws, keys):
    wanted = {norm(k) for k in keys}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
        cols = {norm(c.value): c.column for c in row if c.value is not None}
        if wanted & set(cols):
            return row[0].row, cols
    return None, {}


def write_sheet(wb, name, rows):
    if not rows:
        return
    keys = list(rows[0].keys())
    for r in rows[1:]:
        for k in r:
            if k not in keys:
                keys.append(k)
    if name in wb.sheetnames:
        ws = wb[name]
    else:
        ws = wb.create_sheet(name)
    hdr_row, cols = find_header(ws, keys)
    if hdr_row is None:
        hdr_row = 1 if ws.max_row == 1 and ws.cell(1, 1).value is None else ws.max_row + 2
        cols = {}
    next_col = max(cols.values(), default=0) + 1
    for k in keys:
        if norm(k) not in cols:
            c = ws.cell(hdr_row, next_col, k)
            c.font = Font(bold=True)
            c.fill = HEADER_FILL
            cols[norm(k)] = next_col
            next_col += 1
    style_row = hdr_row + 1
    styles = {col: copy.copy(ws.cell(style_row, col)._style) for col in cols.values()}
    has_style = any(ws.cell(style_row, col).has_style for col in cols.values())
    for i, row in enumerate(rows):
        r = hdr_row + 1 + i
        for k, v in row.items():
            if isinstance(v, (list, tuple)):
                v = ", ".join(str(x) for x in v)
            cell = ws.cell(r, cols[norm(k)], v)
            if has_style:
                cell._style = copy.copy(styles[cols[norm(k)]])
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    last_col = get_column_letter(max(cols.values()))
    ws.auto_filter.ref = f"A{hdr_row}:{last_col}{hdr_row + len(rows)}"
    ws.freeze_panes = ws.cell(hdr_row + 1, 1)
    for k, col in cols.items():
        letter = get_column_letter(col)
        if ws.column_dimensions[letter].width in (None, 13.0):
            longest = max([len(k)] + [len(str(r.get(kk, "") or "")) for r in rows for kk in r if norm(kk) == k])
            ws.column_dimensions[letter].width = min(max(12, longest + 2), 60)


def fill_tokens(wb, meta):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "{{" in c.value:
                    c.value = TOKEN.sub(lambda m: str(meta.get(m.group(1), m.group(0))), c.value)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("json")
    ap.add_argument("output")
    ap.add_argument("--template")
    a = ap.parse_args()
    with open(a.json, encoding="utf-8") as f:
        data = json.load(f)
    wb = openpyxl.load_workbook(a.template) if a.template else openpyxl.Workbook()
    if not a.template:
        wb.remove(wb.active)
    meta = {"DATE": datetime.date.today().isoformat()}
    meta.update({k.upper(): v for k, v in data.get("meta", {}).items()})
    for name, rows in data.get("sheets", {}).items():
        write_sheet(wb, name, rows)
    fill_tokens(wb, meta)
    wb.save(a.output)
    print(f"Wrote {a.output}: " + ", ".join(f"{n}={len(r)} rows" for n, r in data.get("sheets", {}).items()))


if __name__ == "__main__":
    main()
