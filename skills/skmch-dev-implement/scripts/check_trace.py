#!/usr/bin/env python3
"""Traceability check: every requirement ID in the source must appear in the targets.

Reads .md, .txt, .sql, .docx and .xlsx files. Example:
  python check_trace.py --source SRS.docx --target TestCases.xlsx
  python check_trace.py --source SRS.md --target Design.md --target db/ --ids FR,NFR

Exit code 1 when any source ID is missing from all targets (withdrawn IDs excluded).
"""
import argparse
import os
import re
import sys

ID_RE = {
    "FR": r"\bFR-\d{3}\b",
    "NFR": r"\bNFR-[A-Z]{3}-\d{2}\b",
    "BR": r"\bBR-\d{2,3}\b",
    "RULE": r"\bRULE-\d{2,3}\b",
    "DS": r"\bDS-\d{2,3}\b",
    "TC": r"\bTC-\d{3}\b",
}


def read_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        return "\n".join(" | ".join(str(v) for v in row if v is not None)
                         for ws in wb.worksheets for row in ws.iter_rows(values_only=True))
    with open(path, encoding="utf-8", errors="ignore") as f:
        return f.read()


def gather(paths):
    text = []
    for p in paths:
        if os.path.isdir(p):
            for root, _, files in os.walk(p):
                for fn in files:
                    if os.path.splitext(fn)[1].lower() in (".md", ".txt", ".sql", ".pks", ".pkb", ".docx", ".xlsx"):
                        text.append(read_text(os.path.join(root, fn)))
        else:
            text.append(read_text(p))
    return "\n".join(text)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--source", action="append", required=True)
    ap.add_argument("--target", action="append", required=True)
    ap.add_argument("--ids", default="FR,NFR", help="comma list of ID kinds: " + ",".join(ID_RE))
    a = ap.parse_args()
    src, tgt = gather(a.source), gather(a.target)
    withdrawn = set(re.findall(r"\b([A-Z]+-(?:[A-Z]{3}-)?\d{2,3})\b[^\n]{0,80}\bWITHDRAWN\b", src))
    missing_any = False
    for kind in [k.strip().upper() for k in a.ids.split(",") if k.strip()]:
        ids = sorted(set(re.findall(ID_RE[kind], src)) - withdrawn)
        found = set(re.findall(ID_RE[kind], tgt))
        missing = [i for i in ids if i not in found]
        print(f"{kind}: {len(ids)} in source, {len(ids) - len(missing)} covered, {len(missing)} missing")
        if missing:
            missing_any = True
            print("  missing: " + ", ".join(missing))
        extra = sorted(found - set(ids) - withdrawn)
        if extra:
            print(f"  in target but not in source (check for typos/stale IDs): {', '.join(extra)}")
    sys.exit(1 if missing_any else 0)


if __name__ == "__main__":
    main()
