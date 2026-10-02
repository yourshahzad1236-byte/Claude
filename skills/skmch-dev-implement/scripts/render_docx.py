#!/usr/bin/env python3
"""Render a Markdown document into a Word (.docx) template.

The template is any .docx that contains:
  * {{PLACEHOLDER}} tokens (cover page, headers, footers, tables) that are filled from --meta
  * one paragraph whose entire text is {{BODY}}. The rendered Markdown replaces it.
    Without one, the content is appended at the end.

Supported Markdown: '#'..'####' headings, paragraphs, '-'/'*' bullets (nested by
2-space indent), '1.' numbered lists, pipe tables, fenced ``` code blocks,
'> ' notes, **bold**, *italic*, `code`, '<!-- pagebreak -->', and
'<!-- highlight -->' ... '<!-- /highlight -->', which renders everything between the
markers (headings, text, tables) shaded and boxed, e.g. the Impact Analysis section.

Usage:
  python render_docx.py content.md out.docx --template templates/srs_template.docx \
      --meta DOC_ID=CR-2026-014 --meta VERSION=0.1 --meta STATUS=DRAFT ...

Needs python-docx (pip install python-docx).
"""
import argparse
import datetime
import re
import sys

try:
    import docx
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
except ImportError:
    sys.exit("python-docx is required: pip install python-docx")

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)")
TOKEN = re.compile(r"\{\{([A-Z0-9_]+)\}\}")
HIGHLIGHT_FILL = "FFF2CC"      # light amber background for highlighted content
HIGHLIGHT_HEAD = "F4B183"      # stronger amber for highlighted table headers
HIGHLIGHT_BORDER = "C55A11"    # dark orange border


def style_or_default(doc, name, fallback="Normal"):
    try:
        doc.styles[name]
        return name
    except KeyError:
        return fallback


def add_inline(paragraph, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def highlight_paragraph(paragraph):
    """Shade a paragraph and draw a box border around it."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "12" if side == "left" else "4")
        b.set(qn("w:space"), "4")
        b.set(qn("w:color"), HIGHLIGHT_BORDER)
        borders.append(b)
    p_pr.append(borders)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), HIGHLIGHT_FILL)
    p_pr.append(shd)


def replace_tokens_in_paragraph(paragraph, meta):
    full = "".join(r.text for r in paragraph.runs)
    if "{{" not in full:
        return
    new = TOKEN.sub(lambda m: meta.get(m.group(1), m.group(0)), full)
    if new == full or not paragraph.runs:
        return
    paragraph.runs[0].text = new
    for run in paragraph.runs[1:]:
        run.text = ""


def iter_paragraphs(container):
    for p in container.paragraphs:
        yield p
    for t in container.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from iter_paragraphs(cell)


def fill_tokens(doc, meta):
    for p in iter_paragraphs(doc):
        replace_tokens_in_paragraph(p, meta)
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            if part is not None:
                for p in iter_paragraphs(part):
                    replace_tokens_in_paragraph(p, meta)


class Renderer:
    def __init__(self, doc, anchor):
        self.doc = doc
        self.anchor = anchor  # paragraph to insert before (None = append)
        self.highlight = False  # inside <!-- highlight --> ... <!-- /highlight -->

    def _place(self, element):
        if self.anchor is not None:
            self.anchor._p.addprevious(element)

    def paragraph(self, text="", style="Normal"):
        p = self.doc.add_paragraph(style=style_or_default(self.doc, style))
        add_inline(p, text)
        if self.highlight and text:
            highlight_paragraph(p)
        self._place(p._p)
        return p

    def heading(self, text, level):
        p = self.doc.add_paragraph(style=style_or_default(self.doc, f"Heading {level}"))
        add_inline(p, text)
        if self.highlight:
            highlight_paragraph(p)
        self._place(p._p)

    def code(self, lines):
        for line in lines or [""]:
            p = self.doc.add_paragraph(style=style_or_default(self.doc, "No Spacing"))
            run = p.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            if self.highlight:
                highlight_paragraph(p)
            self._place(p._p)

    def note(self, text):
        p = self.paragraph("", "Normal")
        run = p.add_run(text)
        run.italic = True
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        if self.highlight:
            highlight_paragraph(p)

    def pagebreak(self):
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
        self._place(p._p)

    def table(self, rows):
        cols = max(len(r) for r in rows)
        t = self.doc.add_table(rows=len(rows), cols=cols)
        t.style = style_or_default(self.doc, "Table Grid", None) or t.style
        for i, row in enumerate(rows):
            for j in range(cols):
                cell = t.cell(i, j)
                cell.text = ""
                add_inline(cell.paragraphs[0], row[j] if j < len(row) else "")
                for run in cell.paragraphs[0].runs:
                    run.font.size = Pt(9)
                    if i == 0:
                        run.bold = True
                if i == 0:
                    shade(cell, HIGHLIGHT_HEAD if self.highlight else "D9E2F3")
                elif self.highlight:
                    shade(cell, HIGHLIGHT_FILL)
        self._place(t._tbl)
        # spacer after table
        self.paragraph("")


def split_row(line):
    line = line.strip().strip("|")
    return [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", line)]


def render(md, r):
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped == "<!-- pagebreak -->":
            r.pagebreak(); i += 1; continue
        if stripped == "<!-- highlight -->":
            r.highlight = True; i += 1; continue
        if stripped == "<!-- /highlight -->":
            r.highlight = False; i += 1; continue
        if stripped.startswith("<!--"):
            # authoring comments are not rendered
            while i < len(lines) and "-->" not in lines[i]:
                i += 1
            i += 1; continue
        if stripped.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i]); i += 1
            r.code(block); i += 1; continue
        m = re.match(r"^(#{1,4})\s+(.*)", stripped)
        if m:
            r.heading(m.group(2).strip(), len(m.group(1))); i += 1; continue
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = split_row(lines[i])
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                r.table(rows)
            continue
        m = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)", line)
        if m:
            depth = len(m.group(1).replace("\t", "  ")) // 2
            numbered = m.group(2)[0].isdigit()
            base = "List Number" if numbered else "List Bullet"
            style = base if depth == 0 else f"{base} {min(depth + 1, 3)}"
            if style_or_default(r.doc, style, None) is None:
                style = base
            r.paragraph(m.group(3), style); i += 1; continue
        if stripped.startswith(">"):
            r.note(stripped.lstrip("> ").strip()); i += 1; continue
        # paragraph: merge consecutive plain lines
        para = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^\s*(#|\||[-*+]\s|\d+[.)]\s|>|```|<!--)", lines[i]):
            para.append(lines[i].strip()); i += 1
        r.paragraph(" ".join(para))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("markdown")
    ap.add_argument("output")
    ap.add_argument("--template", help=".docx template containing {{BODY}} and {{TOKENS}}")
    ap.add_argument("--meta", action="append", default=[], help="KEY=VALUE to fill {{KEY}}")
    a = ap.parse_args()

    meta = {"DATE": datetime.date.today().isoformat()}
    for kv in a.meta:
        k, _, v = kv.partition("=")
        meta[k.strip().upper()] = v
    doc = docx.Document(a.template) if a.template else docx.Document()

    anchor = None
    for p in doc.paragraphs:
        if p.text.strip() == "{{BODY}}":
            anchor = p
            break
    with open(a.markdown, encoding="utf-8") as f:
        render(f.read(), Renderer(doc, anchor))
    if anchor is not None:
        anchor._p.getparent().remove(anchor._p)
    fill_tokens(doc, meta)
    doc.save(a.output)
    left = sorted({m for p in iter_paragraphs(doc) for m in TOKEN.findall(p.text)})
    print(f"Wrote {a.output}")
    if left:
        print("Unfilled placeholders (pass with --meta):", ", ".join(left))


if __name__ == "__main__":
    main()
