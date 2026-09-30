#!/usr/bin/env python3
"""Generate the PLACEHOLDER Word/Excel templates used by the SKMCH SDLC skills.

These stand in until the hospital's official templates are supplied. To use a real
template, save it over the file in skills/<skill>/templates/ (keep the file name).
Put {{BODY}} on its own line where the generated content should go, and use
{{DOC_ID}}, {{VERSION}}, {{STATUS}}, {{DATE}}, {{AUTHOR}} and {{TITLE}} wherever those
values belong. Run this script again only to reset the placeholders.

Usage: python scripts/build_templates.py
"""
import os

import docx
import openpyxl
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "skills")
ORG = "Shaukat Khanum Memorial Cancer Hospital & Research Centre"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)


def doc_template(path, doc_type, extra_approvers):
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for lvl, size in ((1, 16), (2, 13), (3, 11.5), (4, 10.5)):
        h = d.styles[f"Heading {lvl}"]
        h.font.color.rgb = NAVY
        h.font.size = Pt(size)

    sec = d.sections[0]
    sec.header.paragraphs[0].text = f"{ORG}  |  {doc_type}  |  {{{{DOC_ID}}}}  v{{{{VERSION}}}}"
    sec.header.paragraphs[0].runs[0].font.size = Pt(8)
    sec.footer.paragraphs[0].text = "CONFIDENTIAL: SKMCH internal use only  |  Status: {{STATUS}}"
    sec.footer.paragraphs[0].runs[0].font.size = Pt(8)

    for text, size, bold in ((ORG, 12, False), ("Information Systems", 11, False), ("", 11, False),
                             (doc_type, 22, True), ("{{TITLE}}", 16, True)):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.font.color.rgb = NAVY

    d.add_paragraph()
    t = d.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in (("Document ID", "{{DOC_ID}}"), ("Version", "{{VERSION}}"), ("Status", "{{STATUS}}"),
                 ("Date", "{{DATE}}"), ("Prepared by", "{{AUTHOR}}"),
                 ("Source documents", "{{SOURCE}}")):
        c = t.add_row().cells
        c[0].text, c[1].text = k, v
        c[0].paragraphs[0].runs[0].bold = True

    d.add_paragraph()
    d.add_paragraph("Version History", style="Heading 2")
    vh = d.add_table(rows=2, cols=4)
    vh.style = "Table Grid"
    for i, h in enumerate(("Version", "Date", "Author", "Change summary")):
        vh.cell(0, i).text = h
        vh.cell(0, i).paragraphs[0].runs[0].bold = True
    for i, v in enumerate(("{{VERSION}}", "{{DATE}}", "{{AUTHOR}}", "{{CHANGE_SUMMARY}}")):
        vh.cell(1, i).text = v

    d.add_paragraph("Approvals", style="Heading 2")
    ap = d.add_table(rows=1, cols=5)
    ap.style = "Table Grid"
    for i, h in enumerate(("Role", "Name", "Decision", "Signature", "Date")):
        ap.cell(0, i).text = h
        ap.cell(0, i).paragraphs[0].runs[0].bold = True
    for role in extra_approvers:
        row = ap.add_row().cells
        row[0].text = role
    note = d.add_paragraph()
    r = note.add_run("Approval is recorded by the named person only. AI-drafted documents remain DRAFT until then.")
    r.italic = True
    r.font.size = Pt(8)

    d.add_page_break()
    d.add_paragraph("{{BODY}}")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    d.save(path)
    print("wrote", os.path.relpath(path))


THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="1F3A5F")
HDR_FONT = Font(bold=True, color="FFFFFF")


def cover_sheet(wb, doc_type):
    ws = wb.active
    ws.title = "Cover"
    ws["A1"] = ORG
    ws["A1"].font = Font(bold=True, size=12, color="1F3A5F")
    ws["A2"] = doc_type
    ws["A2"].font = Font(bold=True, size=16, color="1F3A5F")
    for i, (k, v) in enumerate((("Document ID", "{{DOC_ID}}"), ("Title", "{{TITLE}}"),
                                ("Version", "{{VERSION}}"), ("Status", "{{STATUS}}"),
                                ("Date", "{{DATE}}"), ("Prepared by", "{{AUTHOR}}"),
                                ("Source SRS", "{{SOURCE}}")), start=4):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 2, v)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 60


def grid_sheet(wb, name, headers, widths):
    ws = wb.create_sheet(name)
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(1, i, h)
        c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        d = ws.cell(2, i)
        d.border = BORDER
        d.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def xlsx_testcases(path):
    wb = openpyxl.Workbook()
    cover_sheet(wb, "Test Case Specification")
    grid_sheet(wb, "Test Cases",
               ["TC ID", "Covers", "Module / Screen", "Title", "Type", "Priority", "Preconditions",
                "Test Data", "Steps", "Expected Result", "Patient-safety critical", "Automatable"],
               [9, 16, 18, 32, 12, 9, 28, 28, 45, 40, 11, 11])
    grid_sheet(wb, "Traceability", ["Requirement ID", "Requirement summary", "Test cases", "Coverage"],
               [16, 50, 30, 12])
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    print("wrote", os.path.relpath(path))


def xlsx_execution(path):
    wb = openpyxl.Workbook()
    cover_sheet(wb, "Test Execution Report")
    grid_sheet(wb, "Execution",
               ["TC ID", "Covers", "Title", "Cycle", "Environment", "Executed by", "Executed on",
                "Actual Result", "Status", "Defect ID", "Evidence / Remarks"],
               [9, 16, 32, 8, 12, 14, 12, 40, 10, 10, 30])
    grid_sheet(wb, "Defects",
               ["Defect ID", "TC ID", "Covers", "Summary", "Severity", "Priority", "Steps to reproduce",
                "Expected", "Actual", "Status", "Assigned to"],
               [9, 9, 14, 36, 10, 9, 40, 28, 28, 10, 14])
    grid_sheet(wb, "Summary", ["Metric", "Value"], [36, 20])
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    print("wrote", os.path.relpath(path))


if __name__ == "__main__":
    j = os.path.join
    doc_template(j(ROOT, "skmch-ba-srs/templates/srs_template.docx"),
                 "Software Requirements Specification",
                 ["Business Analyst (author)", "Solution Architect", "Business Owner / Department Head"])
    doc_template(j(ROOT, "skmch-sa-srs-review/templates/srs_review_template.docx"),
                 "SRS Review Report", ["Solution Architect (reviewer)"])
    doc_template(j(ROOT, "skmch-sa-design-rfc/templates/design_rfc_template.docx"),
                 "Technical Design Document (RFC)",
                 ["Solution Architect (author)", "Lead Architect", "DBA", "Development Lead"])
    doc_template(j(ROOT, "skmch-dev-implement/templates/impl_notes_template.docx"),
                 "Implementation Notes", ["Developer (author)", "Code Reviewer"])
    doc_template(j(ROOT, "skmch-sysdoc-update/templates/sysdoc_update_template.docx"),
                 "System Documentation Update",
                 ["Implementation Engineer (author)", "Solution Architect"])
    xlsx_testcases(j(ROOT, "skmch-qa-testcases/templates/testcases_template.xlsx"))
    xlsx_execution(j(ROOT, "skmch-qa-execute/templates/test_execution_template.xlsx"))
